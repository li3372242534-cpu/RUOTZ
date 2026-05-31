# -*- coding: utf-8 -*-
"""
Build the presentation:  像差的衍射理论 (Diffraction Theory of Aberrations)
Course: 《光学原理》
- Part 1: basic principles (formulas + illustrative figures)
- Part 2: two frontier papers reviewed like a literature talk (figures from the papers)
"""
import os
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

BASE = "/projects/sandbox/ppt_aberration"
GEN = f"{BASE}/figs/gen"
PA = f"{BASE}/figs/paperA"
PB = f"{BASE}/figs/paperB"

INK   = RGBColor(0x1B, 0x2A, 0x4A)   # deep navy
ACC   = RGBColor(0xC0, 0x39, 0x2B)   # accent red
ACC2  = RGBColor(0x1F, 0x6F, 0xB4)   # accent blue
GREY  = RGBColor(0x55, 0x55, 0x55)
LIGHT = RGBColor(0xF3, 0xF4, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x2E, 0x8B, 0x57)

CJK = "Microsoft YaHei"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def set_font(run, size=18, bold=False, color=INK, name=CJK, italic=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    f.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", name)


def rect(slide, l, t, w, h, fill, line=None):
    sp = slide.shapes.add_shape(1, l, t, w, h)  # rectangle
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(1)
    sp.shadow.inherit = False
    return sp


def textbox(slide, l, t, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(0.05); tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.03); tf.margin_bottom = Inches(0.03)
    return tb, tf


def add_para(tf, parts, first=False, align=PP_ALIGN.LEFT, space_after=6, level=0, line=1.05):
    """parts: list of (text, dict-of-font-kwargs)."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.level = level
    try:
        p.line_spacing = line
    except Exception:
        pass
    for txt, kw in parts:
        r = p.add_run(); r.text = txt
        set_font(r, **kw)
    return p


def add_image_fit(slide, path, l, t, w, h, align="center", valign="middle"):
    iw, ih = Image.open(path).size
    box_r = w / h
    img_r = iw / ih
    if img_r > box_r:
        nw = w; nh = int(w / img_r)
    else:
        nh = h; nw = int(h * img_r)
    if align == "center":
        nl = l + (w - nw) // 2
    elif align == "left":
        nl = l
    else:
        nl = l + (w - nw)
    if valign == "middle":
        nt = t + (h - nh) // 2
    elif valign == "top":
        nt = t
    else:
        nt = t + (h - nh)
    slide.shapes.add_picture(path, nl, nt, nw, nh)
    return nl, nt, nw, nh


def header(slide, title, kicker=None):
    rect(slide, 0, 0, SW, Inches(0.95), INK)
    rect(slide, 0, Inches(0.95), SW, Inches(0.06), ACC)
    tb, tf = textbox(slide, Inches(0.5), Inches(0.12), Inches(12.3), Inches(0.78),
                     anchor=MSO_ANCHOR.MIDDLE)
    if kicker:
        add_para(tf, [(kicker + "   ", dict(size=12, bold=True, color=RGBColor(0xAF, 0xC6, 0xE6)))],
                 first=True, space_after=0)
        add_para(tf, [(title, dict(size=24, bold=True, color=WHITE))], space_after=0)
    else:
        add_para(tf, [(title, dict(size=26, bold=True, color=WHITE))], first=True, space_after=0)


def caption(slide, l, t, w, txt):
    tb, tf = textbox(slide, l, t, w, Inches(0.4))
    add_para(tf, [(txt, dict(size=10.5, italic=True, color=GREY))], first=True, align=PP_ALIGN.CENTER, space_after=0)


def footer(slide, n):
    tb, tf = textbox(slide, Inches(0.4), Inches(7.08), Inches(9), Inches(0.35))
    add_para(tf, [("《光学原理》课程报告 · 像差的衍射理论", dict(size=9, color=GREY))], first=True, space_after=0)
    tb2, tf2 = textbox(slide, Inches(12.2), Inches(7.08), Inches(0.9), Inches(0.35))
    add_para(tf2, [(str(n), dict(size=9, color=GREY))], first=True, align=PP_ALIGN.RIGHT, space_after=0)


PAGE = [0]
def newslide(title=None, kicker=None):
    s = prs.slides.add_slide(BLANK)
    if title:
        header(s, title, kicker)
        PAGE[0] += 1
        footer(s, PAGE[0])
    return s


# ----------------------------------------------------------------------------
# 1. TITLE
# ----------------------------------------------------------------------------
s = newslide()
rect(s, 0, 0, SW, SH, INK)
rect(s, 0, Inches(4.55), SW, Inches(0.05), ACC)
tb, tf = textbox(s, Inches(0.9), Inches(1.7), Inches(11.5), Inches(2.6))
add_para(tf, [("像差的衍射理论", dict(size=52, bold=True, color=WHITE))], first=True, space_after=8)
add_para(tf, [("Diffraction Theory of Aberrations", dict(size=24, color=RGBColor(0xAF,0xC6,0xE6)))], space_after=0)
tb2, tf2 = textbox(s, Inches(0.9), Inches(4.75), Inches(11.5), Inches(2.2))
add_para(tf2, [("基本原理   ·   点扩散函数 (PSF)   ·   Zernike 多项式   ·   Nijboer–Zernike 理论", dict(size=16, color=RGBColor(0xDD,0xE4,0xF0)))], first=True, space_after=6)
add_para(tf2, [("前沿文献：① Nijboer–Zernike 理论用于空间引力波探测   ② 深度学习像差反演", dict(size=16, color=RGBColor(0xDD,0xE4,0xF0)))], space_after=14)
add_para(tf2, [("《光学原理》课程报告", dict(size=15, bold=True, color=WHITE))], space_after=2)

# ----------------------------------------------------------------------------
# 2. OUTLINE
# ----------------------------------------------------------------------------
s = newslide("目录 / Outline", "报告结构")
tb, tf = textbox(s, Inches(0.7), Inches(1.4), Inches(6.0), Inches(5.4))
items_l = [
    ("第一部分  基本原理", True),
    ("1. 从几何像差到衍射像差：为什么需要衍射理论", False),
    ("2. 波像差函数 W 与光瞳函数", False),
    ("3. 衍射积分与点扩散函数 (PSF)", False),
    ("4. 理想 PSF：艾里斑 (Airy disk)", False),
    ("5. 各类像差对 PSF 的影响", False),
    ("6. Zernike 圆多项式", False),
    ("7. Nijboer–Zernike 衍射理论", False),
    ("8. 斯特列尔比与 Maréchal 判据", False),
    ("9. 光学传递函数 MTF", False),
    ("10. 扩展 Nijboer–Zernike (ENZ) 理论", False),
]
for i, (t, head) in enumerate(items_l):
    add_para(tf, [(t, dict(size=(18 if head else 14.5), bold=head, color=(ACC if head else INK)))],
             first=(i == 0), space_after=(8 if head else 4))

tb2, tf2 = textbox(s, Inches(7.0), Inches(1.4), Inches(5.7), Inches(5.4))
items_r = [
    ("第二部分  前沿文献调研", True),
    ("文献 A (2026)：将 Nijboer–Zernike 理论用于", False),
    ("空间引力波探测中远场波前误差与 TTL 噪声建模", False),
    ("文献 B (2024)：用卷积神经网络从离焦 PSF 图像", False),
    ("直接预测 Zernike 系数，实现快速像差反演", False),
    ("", False),
    ("第三部分  总结与展望", True),
    ("· 经典衍射像差理论今天依然“活着”", False),
    ("· 正在与计算成像 / 人工智能深度融合", False),
]
for i, (t, head) in enumerate(items_r):
    add_para(tf2, [(t, dict(size=(18 if head else 14.5), bold=head, color=(ACC if head else INK)))],
             first=(i == 0), space_after=(8 if head else 4))

# ----------------------------------------------------------------------------
# 3. 几何像差 vs 衍射像差
# ----------------------------------------------------------------------------
s = newslide("从几何像差到衍射像差", "第一部分 · 基本原理 (1/10)")
tb, tf = textbox(s, Inches(0.7), Inches(1.35), Inches(11.9), Inches(5.4))
add_para(tf, [("几何光学", dict(size=18, bold=True, color=ACC2)),
              ("：用“光线”描述成像，像差 = 光线偏离理想像点的程度（横向像差、纵向像差）。",
               dict(size=15))], first=True, space_after=8)
add_para(tf, [("• 适用于像差较大、远大于波长的系统；可用 Seidel 五种初级像差（球差、彗差、像散、场曲、畸变）描述。",
               dict(size=14))], space_after=10)
add_para(tf, [("衍射光学", dict(size=18, bold=True, color=ACC)),
              ("：当系统已被校正到接近“衍射极限”，残余像差与波长 λ 同量级，光线模型失效。",
               dict(size=15))], space_after=8)
add_para(tf, [("• 此时必须把光当作", dict(size=14)),
              ("波", dict(size=14, bold=True, color=ACC)),
              (" 来处理：像点不再是一个点，而是一个由衍射决定的", dict(size=14)),
              ("强度分布（点扩散函数 PSF）", dict(size=14, bold=True, color=ACC)),
              ("。", dict(size=14))], space_after=8)
add_para(tf, [("• 像差的作用：改变光瞳上的相位分布 → 改变 PSF 的形状 → 降低成像分辨率与对比度。",
               dict(size=14))], space_after=14)
rect(s, Inches(0.7), Inches(5.55), Inches(11.9), Inches(1.15), LIGHT)
tb2, tf2 = textbox(s, Inches(0.95), Inches(5.62), Inches(11.4), Inches(1.0), anchor=MSO_ANCHOR.MIDDLE)
add_para(tf2, [("核心思想：", dict(size=15, bold=True, color=ACC)),
               ("“像差的衍射理论” 研究 ", dict(size=14)),
               ("波像差 W(x,y)", dict(size=14, bold=True, color=INK)),
               (" 如何通过 ", dict(size=14)),
               ("衍射积分", dict(size=14, bold=True, color=INK)),
               (" 决定像面光场（PSF / OTF），并据此评价与反演光学系统。", dict(size=14))],
         first=True, space_after=0)

# ----------------------------------------------------------------------------
# 4. 波像差函数
# ----------------------------------------------------------------------------
s = newslide("波像差函数 W 与光瞳函数", "第一部分 · 基本原理 (2/10)")
tb, tf = textbox(s, Inches(0.65), Inches(1.3), Inches(6.4), Inches(5.5))
add_para(tf, [("波像差 W：", dict(size=17, bold=True, color=ACC)),
              ("出射光瞳上，实际波前相对理想参考球面（球心在高斯像点）的光程差。",
               dict(size=14))], first=True, space_after=8)
add_para(tf, [("相位形式：  Φ(ρ,θ) = (2π/λ) · W(ρ,θ)", dict(size=15, bold=True, color=INK, name="Cambria Math"))], space_after=10)
add_para(tf, [("光瞳函数（含像差）：", dict(size=15, bold=True, color=ACC2))], space_after=4)
add_para(tf, [("P(ρ,θ) = A(ρ,θ) · exp[ i · (2π/λ) W(ρ,θ) ]", dict(size=15, bold=True, color=INK, name="Cambria Math"))], space_after=10)
add_para(tf, [("• A(ρ,θ)：振幅（光阑形状 / 透过率）", dict(size=13.5))], space_after=3)
add_para(tf, [("• 相位项 exp(iΦ)：像差全部信息所在", dict(size=13.5))], space_after=3)
add_para(tf, [("• ρ：归一化光瞳半径 (0~1)，θ：方位角", dict(size=13.5))], space_after=10)
add_para(tf, [("一句话：", dict(size=14, bold=True, color=ACC)),
              ("像差 = 光瞳上的相位畸变。", dict(size=14))], space_after=0)
add_image_fit(s, f"{GEN}/01_wavefront_schematic.png", Inches(7.0), Inches(1.5), Inches(5.9), Inches(4.6))
caption(s, Inches(7.0), Inches(6.15), Inches(5.9), "示意图（自绘）：波像差 W 即实际波前对理想参考球面的偏离")

# ----------------------------------------------------------------------------
# 5. 衍射积分与 PSF
# ----------------------------------------------------------------------------
s = newslide("衍射积分与点扩散函数 (PSF)", "第一部分 · 基本原理 (3/10)")
tb, tf = textbox(s, Inches(0.7), Inches(1.35), Inches(11.9), Inches(5.4))
add_para(tf, [("像面复振幅 = 光瞳函数的（夫琅禾费 / Debye）衍射积分：", dict(size=15, bold=True, color=ACC))], first=True, space_after=8)
add_para(tf, [("U(v,φ;f) = ∬  exp[ iΦ(ρ,θ) ] · exp[ i·½f·ρ² ] · exp[ −i·v·ρ·cos(θ−φ) ] · ρ dρ dθ",
               dict(size=15.5, bold=True, color=INK, name="Cambria Math"))], align=PP_ALIGN.CENTER, space_after=10)
add_para(tf, [("点扩散函数（PSF，像面强度分布）：   ", dict(size=15, bold=True, color=ACC2)),
              ("I(v,φ) = | U(v,φ) |²", dict(size=15.5, bold=True, color=INK, name="Cambria Math"))], space_after=12)
add_para(tf, [("符号说明：", dict(size=14, bold=True, color=INK))], space_after=4)
for t in [
    "v = (2π/λ)(a/R)·r   —— 像面归一化径向坐标（a 为光瞳半径，R 为参考球面半径，r 为实际像面半径）",
    "f —— 归一化离焦参量（轴向坐标）；φ —— 像面方位角",
    "Φ(ρ,θ) = (2π/λ)W —— 像差相位；积分在单位圆光瞳上进行",
]:
    add_para(tf, [("• " + t, dict(size=13.5))], space_after=4)
rect(s, Inches(0.7), Inches(5.7), Inches(11.9), Inches(1.0), LIGHT)
tb2, tf2 = textbox(s, Inches(0.95), Inches(5.76), Inches(11.4), Inches(0.88), anchor=MSO_ANCHOR.MIDDLE)
add_para(tf2, [("数学本质：", dict(size=15, bold=True, color=ACC)),
               ("PSF 是光瞳函数的傅里叶变换的模平方；OTF（光学传递函数）是光瞳函数的自相关。像差 → 相位畸变 → PSF 展宽、MTF 下降。",
                dict(size=14))], first=True, space_after=0)

# ----------------------------------------------------------------------------
# 6. Airy
# ----------------------------------------------------------------------------
s = newslide("理想 PSF：艾里斑 (Airy disk)", "第一部分 · 基本原理 (4/10)")
add_image_fit(s, f"{GEN}/03_airy.png", Inches(0.6), Inches(1.45), Inches(7.5), Inches(5.0), align="left")
tb, tf = textbox(s, Inches(8.3), Inches(1.6), Inches(4.6), Inches(5.0))
add_para(tf, [("无像差、圆形光阑下，PSF 为著名的艾里斑：", dict(size=14, bold=True, color=ACC))], first=True, space_after=8)
add_para(tf, [("I(v) = [ 2J₁(v) / v ]²", dict(size=18, bold=True, color=INK, name="Cambria Math"))], align=PP_ALIGN.CENTER, space_after=10)
add_para(tf, [("• J₁ 为一阶贝塞尔函数", dict(size=13.5))], space_after=4)
add_para(tf, [("• 第一暗环位于 v = 3.832", dict(size=13.5))], space_after=4)
add_para(tf, [("• 中央亮斑集中约 84% 能量", dict(size=13.5))], space_after=10)
add_para(tf, [("艾里斑大小 → 衍射极限分辨率（瑞利判据）。这是衡量一切像差影响的“基准”。",
               dict(size=14))], space_after=0)
caption(s, Inches(0.6), Inches(6.4), Inches(7.5), "自绘（由光瞳函数 FFT 计算）：艾里斑二维分布与径向强度")

# ----------------------------------------------------------------------------
# 7. 各类像差对 PSF 的影响
# ----------------------------------------------------------------------------
s = newslide("各类像差对 PSF 的影响", "第一部分 · 基本原理 (5/10)")
add_image_fit(s, f"{GEN}/04_aberrated_psf.png", Inches(0.5), Inches(1.5), Inches(12.3), Inches(3.4))
caption(s, Inches(0.5), Inches(4.85), Inches(12.3), "自绘（FFT 计算，各像差 RMS=0.25λ）：像差越大，PSF 偏离艾里斑越明显")
tb, tf = textbox(s, Inches(0.7), Inches(5.3), Inches(11.9), Inches(1.7))
add_para(tf, [("• 离焦 (Defocus)：", dict(size=13.5, bold=True, color=ACC2)),
              ("亮斑展宽、能量向外环转移；   ", dict(size=13.5)),
              ("• 像散 (Astigmatism)：", dict(size=13.5, bold=True, color=ACC2)),
              ("沿两正交方向拉伸成椭圆/十字。", dict(size=13.5))], first=True, space_after=6)
add_para(tf, [("• 彗差 (Coma)：", dict(size=13.5, bold=True, color=ACC)),
              ("不对称的“彗星尾”，破坏轴对称；   ", dict(size=13.5)),
              ("• 球差 (Spherical)：", dict(size=13.5, bold=True, color=ACC)),
              ("中心亮斑变弱、外围出现明显光晕。", dict(size=13.5))], space_after=6)
add_para(tf, [("结论：", dict(size=14, bold=True, color=INK)),
              ("不同像差在像面上留下各自“指纹”——这正是后面“由 PSF 反演像差”的物理基础。", dict(size=13.5))], space_after=0)

# ----------------------------------------------------------------------------
# 8. Zernike
# ----------------------------------------------------------------------------
s = newslide("Zernike 圆多项式", "第一部分 · 基本原理 (6/10)")
add_image_fit(s, f"{GEN}/02_zernike_maps.png", Inches(7.0), Inches(1.45), Inches(6.0), Inches(5.0))
caption(s, Inches(7.0), Inches(6.4), Inches(6.0), "自绘：前几项 Zernike 多项式的相位分布")
tb, tf = textbox(s, Inches(0.65), Inches(1.4), Inches(6.2), Inches(5.5))
add_para(tf, [("把波像差展开为一组在单位圆上正交的多项式：", dict(size=14, bold=True, color=ACC))], first=True, space_after=8)
add_para(tf, [("W(ρ,θ) = Σ cₙᵐ · Zₙᵐ(ρ,θ)", dict(size=17, bold=True, color=INK, name="Cambria Math"))], align=PP_ALIGN.CENTER, space_after=6)
add_para(tf, [("Zₙᵐ(ρ,θ) = Rₙᵐ(ρ) · {cos / sin}(mθ)", dict(size=14, bold=True, color=INK, name="Cambria Math"))], align=PP_ALIGN.CENTER, space_after=10)
add_para(tf, [("为什么用它？三大优点：", dict(size=14, bold=True, color=ACC2))], space_after=4)
add_para(tf, [("① 正交性 → 总方差 = 各项系数平方和，", dict(size=13.5)),
              ("σ²_W = Σ (cₙᵐ)²", dict(size=13.5, bold=True, name="Cambria Math"))], space_after=4)
add_para(tf, [("② 每一项对应一种“平衡像差”（如 Z₄⁰ = 球差与离焦的最优平衡）", dict(size=13.5))], space_after=4)
add_para(tf, [("③ 与经典 Seidel 像差一一对应：", dict(size=13.5))], space_after=2)
add_para(tf, [("   倾斜 Z₁¹ · 离焦 Z₂⁰ · 像散 Z₂² · 彗差 Z₃¹ · 球差 Z₄⁰", dict(size=13, bold=True, color=INK, name="Cambria Math"))], space_after=0)

# ----------------------------------------------------------------------------
# 9. Nijboer-Zernike
# ----------------------------------------------------------------------------
s = newslide("Nijboer–Zernike 衍射理论", "第一部分 · 基本原理 (7/10)")
tb, tf = textbox(s, Inches(0.65), Inches(1.35), Inches(6.4), Inches(5.5))
add_para(tf, [("Nijboer (1942–47) & Zernike：", dict(size=15, bold=True, color=ACC)),
              ("把像差用 Zernike 多项式展开后，", dict(size=14)),
              ("逐项解析地", dict(size=14, bold=True, color=ACC)),
              ("计算其衍射 PSF。", dict(size=14))], first=True, space_after=8)
add_para(tf, [("关键结果——每个 Zernike 项的衍射积分都是一个解析的贝塞尔函数：", dict(size=14, bold=True, color=ACC2))], space_after=6)
add_para(tf, [("∫₀¹ Rₙᵐ(ρ) Jₘ(vρ) ρ dρ = (−1)^((n−m)/2) · Jₙ₊₁(v)/v",
               dict(size=15, bold=True, color=INK, name="Cambria Math"))], align=PP_ALIGN.CENTER, space_after=10)
add_para(tf, [("意义：", dict(size=14, bold=True, color=INK))], space_after=3)
add_para(tf, [("• 把“算衍射积分”变成“查贝塞尔函数”，给出小像差下 PSF 的解析表达；", dict(size=13.5))], space_after=4)
add_para(tf, [("• 是连接 波像差(Zernike) ↔ 像面强度(PSF) 的桥梁；", dict(size=13.5))], space_after=4)
add_para(tf, [("• 历史局限：原始理论仅适用于小像差、无离焦的标量近似。", dict(size=13.5, color=ACC))], space_after=0)
add_image_fit(s, f"{GEN}/07_nijboer_zernike.png", Inches(7.1), Inches(1.55), Inches(5.8), Inches(4.7))
caption(s, Inches(7.1), Inches(6.3), Inches(5.8), "自绘：各 Zernike 项衍射积分 = 解析贝塞尔函数 Jₙ₊₁(v)/v")

# ----------------------------------------------------------------------------
# 10. Strehl
# ----------------------------------------------------------------------------
s = newslide("斯特列尔比与 Maréchal 判据", "第一部分 · 基本原理 (8/10)")
add_image_fit(s, f"{GEN}/05_strehl.png", Inches(0.6), Inches(1.5), Inches(7.2), Inches(5.0), align="left")
caption(s, Inches(0.6), Inches(6.4), Inches(7.2), "自绘：Strehl 比随 RMS 波像差的变化")
tb, tf = textbox(s, Inches(8.0), Inches(1.6), Inches(4.9), Inches(5.0))
add_para(tf, [("斯特列尔比 S：", dict(size=16, bold=True, color=ACC)),
              ("有像差 PSF 的峰值强度 / 无像差时的峰值强度。", dict(size=13.5))], first=True, space_after=8)
add_para(tf, [("S = | ⟨ exp(iΦ) ⟩ |²", dict(size=16, bold=True, color=INK, name="Cambria Math"))], align=PP_ALIGN.CENTER, space_after=8)
add_para(tf, [("小像差近似（Maréchal）：", dict(size=14, bold=True, color=ACC2))], space_after=4)
add_para(tf, [("S ≈ exp[ −(2πσ_W)² ] ≈ 1 − (2πσ_W)²", dict(size=14.5, bold=True, color=INK, name="Cambria Math"))], align=PP_ALIGN.CENTER, space_after=10)
add_para(tf, [("Maréchal 判据（衍射极限）：", dict(size=14, bold=True, color=GREEN))], space_after=4)
add_para(tf, [("S ≥ 0.8  ⟺  RMS 波像差 σ_W ≤ λ/14", dict(size=14.5, bold=True, color=GREEN, name="Cambria Math"))], space_after=10)
add_para(tf, [("即：只要 RMS 波前误差不超过约 λ/14，系统就可视为“衍射极限”成像。", dict(size=13.5))], space_after=0)

# ----------------------------------------------------------------------------
# 11. MTF
# ----------------------------------------------------------------------------
s = newslide("光学传递函数 (MTF)", "第一部分 · 基本原理 (9/10)")
add_image_fit(s, f"{GEN}/06_mtf.png", Inches(0.6), Inches(1.5), Inches(7.2), Inches(5.0), align="left")
caption(s, Inches(0.6), Inches(6.4), Inches(7.2), "自绘：无像差 vs 含彗差/球差时的 MTF 曲线")
tb, tf = textbox(s, Inches(8.0), Inches(1.6), Inches(4.9), Inches(5.0))
add_para(tf, [("OTF = 光瞳函数的自相关；MTF = |OTF|。", dict(size=14, bold=True, color=ACC))], first=True, space_after=8)
add_para(tf, [("MTF(ν) 描述系统对不同空间频率 ν 的对比度传递能力。", dict(size=13.5))], space_after=8)
add_para(tf, [("无像差（衍射极限）MTF：", dict(size=13.5, bold=True, color=ACC2))], space_after=3)
add_para(tf, [("(2/π)[ arccos(ν) − ν√(1−ν²) ]", dict(size=14, bold=True, color=INK, name="Cambria Math"))], align=PP_ALIGN.CENTER, space_after=10)
add_para(tf, [("像差的作用：", dict(size=13.5, bold=True, color=ACC))], space_after=3)
add_para(tf, [("• 中、高频对比度下降 → 图像“变糊”、细节丢失；", dict(size=13.5))], space_after=4)
add_para(tf, [("• 截止频率 ν_c 不变，但曲线整体被压低。", dict(size=13.5))], space_after=8)
add_para(tf, [("PSF 与 MTF 互为傅里叶变换——同一信息的两种视角。", dict(size=13.5, italic=True, color=GREY))], space_after=0)

# ----------------------------------------------------------------------------
# 12. ENZ
# ----------------------------------------------------------------------------
s = newslide("扩展 Nijboer–Zernike (ENZ) 理论", "第一部分 · 基本原理 (10/10)")
add_image_fit(s, f"{GEN}/08_through_focus.png", Inches(0.5), Inches(1.5), Inches(12.3), Inches(3.3))
caption(s, Inches(0.5), Inches(4.8), Inches(12.3), "自绘：含球差时通过焦点的 PSF 演化——ENZ 可解析描述并用于像差反演")
tb, tf = textbox(s, Inches(0.7), Inches(5.25), Inches(11.9), Inches(1.8))
add_para(tf, [("Janssen, Braat 等 (2002 起)：", dict(size=14, bold=True, color=ACC)),
              ("把经典 N–Z 理论从“小像差、无离焦”推广到现代实用场景：", dict(size=13.5))], first=True, space_after=6)
add_para(tf, [("① 任意大小的像差与离焦；② 复光瞳函数（同时含相位 + 振幅/透过率变化）；③ 高数值孔径下的矢量衍射。",
               dict(size=13.5))], space_after=6)
add_para(tf, [("关键价值：", dict(size=14, bold=True, color=ACC2)),
              ("用 Bessel 级数 Vₙᵐ(r,f) 高效、精确地计算 PSF，并可求解", dict(size=13.5)),
              ("反问题——从实测的通过焦点强度图像反演出复光瞳（像差）", dict(size=13.5, bold=True, color=ACC)),
              ("。这正是文献 A、B 的共同根基。", dict(size=13.5))], space_after=0)

# ----------------------------------------------------------------------------
# 13. PART 2 divider
# ----------------------------------------------------------------------------
s = newslide()
rect(s, 0, 0, SW, SH, INK)
rect(s, Inches(0.9), Inches(2.7), Inches(0.18), Inches(2.1), ACC)
tb, tf = textbox(s, Inches(1.3), Inches(2.6), Inches(11), Inches(2.4))
add_para(tf, [("第二部分", dict(size=22, color=RGBColor(0xAF,0xC6,0xE6)))], first=True, space_after=6)
add_para(tf, [("前沿文献调研", dict(size=46, bold=True, color=WHITE))], space_after=16)
add_para(tf, [("两篇近年文献，分别展示经典衍射像差理论在", dict(size=17, color=RGBColor(0xDD,0xE4,0xF0)))], space_after=4)
add_para(tf, [("空间引力波探测", dict(size=17, bold=True, color=WHITE)),
              (" 与 ", dict(size=17, color=RGBColor(0xDD,0xE4,0xF0))),
              ("人工智能像差反演", dict(size=17, bold=True, color=WHITE)),
              (" 中的最新应用。", dict(size=17, color=RGBColor(0xDD,0xE4,0xF0)))], space_after=0)

# ----------------------------------------------------------------------------
# 14. Paper A overview
# ----------------------------------------------------------------------------
s = newslide("文献 A · 概览", "前沿文献 ①")
tb, tf = textbox(s, Inches(0.7), Inches(1.25), Inches(11.9), Inches(1.4))
add_para(tf, [("Analytical Modeling of Far-Field Wavefront Error with Beam-Waist and "
               "Lateral-Shift Effects in Spaceborne Laser Interferometry",
               dict(size=16, bold=True, color=INK))], first=True, space_after=4)
add_para(tf, [("Ya-Zheng Tao, Rui-Hong Gao, Guangzhou Xu, Yue-Liang Wu — arXiv:2604.26371 (2026), CC BY 4.0",
               dict(size=12, italic=True, color=GREY))], space_after=0)
tb2, tf2 = textbox(s, Inches(0.7), Inches(2.75), Inches(11.9), Inches(4.2))
add_para(tf2, [("研究背景：", dict(size=15, bold=True, color=ACC)),
               ("空间引力波探测（如 LISA / 太极 Taiji）卫星间相距数百万公里，靠激光干涉测距，"
                "精度需达皮米 (pm) 量级。", dict(size=14))], first=True, space_after=8)
add_para(tf2, [("核心问题——TTL 噪声：", dict(size=15, bold=True, color=ACC2)),
               ("远场波前误差 (WFE) 与激光指向抖动耦合，产生“倾斜-平移”(tilt-to-length) 噪声，"
                "是主要误差源之一。", dict(size=14))], space_after=8)
add_para(tf2, [("本文做法：", dict(size=15, bold=True, color=GREEN)),
               ("把经典 Nijboer–Zernike 模型 ", dict(size=14)),
               ("扩展", dict(size=14, bold=True, color=ACC)),
               ("到截断高斯光束的远场 WFE，引入两个实用初始参数：", dict(size=14))], space_after=6)
add_para(tf2, [("   • 束腰/孔径比  q = w₀ / rₐ   （光束截断程度）", dict(size=14, bold=True, color=INK, name="Cambria Math"))], space_after=4)
add_para(tf2, [("   • 归一化横向光斑偏移比  s_r           （对准误差）", dict(size=14, bold=True, color=INK, name="Cambria Math"))], space_after=10)
add_para(tf2, [("意义：", dict(size=15, bold=True, color=ACC)),
               ("说明 80 年前的衍射像差理论，如今仍是空间引力波探测光学设计的解析工具。",
                dict(size=14))], space_after=0)

# ----------------------------------------------------------------------------
# 15. Paper A - figures (power + WFE)
# ----------------------------------------------------------------------------
s = newslide("文献 A · 束腰/孔径比 q 的影响", "前沿文献 ①")
add_image_fit(s, f"{PA}/x3.png", Inches(0.6), Inches(1.45), Inches(6.0), Inches(3.9), align="center")
caption(s, Inches(0.6), Inches(5.35), Inches(6.0), "原文 Fig.2：远端接收功率随 q 变化，峰值在 q≈0.892")
add_image_fit(s, f"{PA}/x4.png", Inches(6.9), Inches(1.45), Inches(6.0), Inches(3.9), align="center")
caption(s, Inches(6.9), Inches(5.35), Inches(6.0), "原文 Fig.3：若干偶数阶 Zernike 项对远场 WFE 的一阶贡献随 q 变化")
tb, tf = textbox(s, Inches(0.7), Inches(5.75), Inches(11.9), Inches(1.3))
add_para(tf, [("• 接收功率", dict(size=13.5, bold=True, color=ACC2)),
              ("在 q≈0.89 处最大——这是传统上选择束腰大小的依据；", dict(size=13.5))], first=True, space_after=5)
add_para(tf, [("• 但作者指出：", dict(size=13.5, bold=True, color=ACC)),
              ("适当减小 q 还能降低远场 WFE。蒙特卡洛模拟显示 q 由 1→0.9、0.9→0.8 分别使平均 WFE 降低约 10% 与 14%。",
               dict(size=13.5))], space_after=5)
add_para(tf, [("→ 给出了“功率—波前误差”之间的权衡，为束参数优化提供解析依据。", dict(size=13.5, italic=True, color=GREY))], space_after=0)

# ----------------------------------------------------------------------------
# 16. Paper A - Monte Carlo
# ----------------------------------------------------------------------------
s = newslide("文献 A · 蒙特卡洛验证与结论", "前沿文献 ①")
add_image_fit(s, f"{PA}/x9.png", Inches(0.7), Inches(1.5), Inches(7.6), Inches(4.3), align="left")
caption(s, Inches(0.7), Inches(5.85), Inches(7.6), "原文 Fig.5：随机初始像差下，不同 q 的远场 WFE 蒙特卡洛对比")
tb, tf = textbox(s, Inches(8.5), Inches(1.6), Inches(4.4), Inches(5.2))
add_para(tf, [("主要结论：", dict(size=16, bold=True, color=ACC))], first=True, space_after=8)
add_para(tf, [("① 减小 q 在牺牲少量功率的同时显著降低 WFE，存在最优折中。", dict(size=13.5))], space_after=8)
add_para(tf, [("② 横向光斑偏移：", dict(size=13.5, bold=True, color=ACC2))], space_after=3)
add_para(tf, [("太极望远镜入瞳偏移 2 μm（s_r=0.001）对应相位耦合系数约 0.0892 pm/nrad，"
               "接近 0.1 pm/nrad 的 TTL 指标。", dict(size=13.5))], space_after=8)
add_para(tf, [("③ 偏移—像差的交叉耦合项很小，工程容差估计中可忽略。", dict(size=13.5))], space_after=8)
add_para(tf, [("启示：", dict(size=14, bold=True, color=GREEN)),
              ("经典衍射像差理论 + Zernike 展开，为未来空间引力波任务的光学容差设计提供了解析框架。",
               dict(size=13.5))], space_after=0)

# ----------------------------------------------------------------------------
# 17. Paper B overview
# ----------------------------------------------------------------------------
s = newslide("文献 B · 概览", "前沿文献 ②")
tb, tf = textbox(s, Inches(0.7), Inches(1.25), Inches(11.9), Inches(1.4))
add_para(tf, [("Direct Zernike Coefficient Prediction from Point Spread Functions and "
               "Extended Images using Deep Learning",
               dict(size=16, bold=True, color=INK))], first=True, space_after=4)
add_para(tf, [("Y. E. Kok, A. Bentley, A. Parkes, A. J. Wright, M. G. Somekh, M. Pound "
               "(Univ. of Nottingham) — arXiv:2404.15231 (2024)",
               dict(size=12, italic=True, color=GREY))], space_after=0)
tb2, tf2 = textbox(s, Inches(0.7), Inches(2.75), Inches(11.9), Inches(4.2))
add_para(tf2, [("研究背景：", dict(size=15, bold=True, color=ACC)),
               ("荧光显微等成像中，系统与样品引入的像差会严重降低分辨率。传统自适应光学 (AO) 依赖"
                "波前传感器 + 变形镜，并需迭代搜索，速度慢、硬件复杂。", dict(size=14))], first=True, space_after=8)
add_para(tf2, [("本文做法：", dict(size=15, bold=True, color=GREEN)),
               ("用卷积神经网络 (ResNet-50)，从 2~3 张 ", dict(size=14)),
               ("相位差异 (phase-diverse)", dict(size=14, bold=True, color=ACC)),
               (" 图像——即焦面上、下、在焦三个位置的 PSF——", dict(size=14)),
               ("直接预测前 25 项 Zernike 系数", dict(size=14, bold=True, color=ACC)),
               ("。", dict(size=14))], space_after=8)
add_para(tf2, [("与经典理论的联系：", dict(size=15, bold=True, color=ACC2)),
               ("这正是 (E)NZ 反问题——“由通过焦点的 PSF 反演光瞳像差”——的数据驱动版本；"
                "网络学到的，是 像差(Zernike) ↔ 衍射 PSF 之间的映射。", dict(size=14))], space_after=8)
add_para(tf2, [("结果：", dict(size=15, bold=True, color=ACC)),
               ("在 60 万张仿真 PSF 上，3 通道输入达到 RMS≈0.10 rad；推广到 2D 扩展样品仍约 0.15 rad，"
                "且只需 1 次（或少数几次迭代）即可校正。", dict(size=14))], space_after=0)

# ----------------------------------------------------------------------------
# 18. Paper B - phase diverse + architecture
# ----------------------------------------------------------------------------
s = newslide("文献 B · 相位差异输入与网络框架", "前沿文献 ②")
# three phase-diverse PSFs
labels = [("psfSNegativeDefocus.png", "负离焦 I₋₁"),
          ("psfSAtFocus.png", "在焦 I₀"),
          ("psfSPositiveFocus.png", "正离焦 I₊₁")]
x0 = Inches(0.6); w = Inches(2.0); gap = Inches(0.15); top = Inches(1.6)
for i, (fn, lab) in enumerate(labels):
    lx = x0 + i * (w + gap)
    add_image_fit(s, f"{PB}/{fn}", lx, top, w, Inches(2.0))
    cb, cf = textbox(s, lx, top + Inches(2.0), w, Inches(0.35))
    add_para(cf, [(lab, dict(size=11, bold=True, color=INK))], first=True, align=PP_ALIGN.CENTER, space_after=0)
caption(s, x0, Inches(4.15), Inches(6.6), "原文 Fig.1：同一像差下，焦面上/下/在焦三张 PSF（提供相位差异信息）")
add_image_fit(s, f"{PB}/architecture.png", Inches(7.3), Inches(1.9), Inches(5.6), Inches(2.0), align="center")
caption(s, Inches(7.3), Inches(3.95), Inches(5.6), "原文 Fig.3：ResNet-50 由 n 张相位差异图像直接回归 Zernike 系数")
tb, tf = textbox(s, Inches(0.7), Inches(4.7), Inches(11.9), Inches(2.2))
add_para(tf, [("• 为什么要“相位差异”？", dict(size=14, bold=True, color=ACC)),
              ("单张在焦 PSF 难以区分正/负号像差；加入离焦图像后，像差符号与大小被唯一编码，反演不再奇异。",
               dict(size=13.5))], first=True, space_after=6)
add_para(tf, [("• 输入：", dict(size=13.5, bold=True, color=ACC2)),
              ("3 通道（−1, 0, +1 离焦）PSF；   ", dict(size=13.5)),
              ("输出：", dict(size=13.5, bold=True, color=ACC2)),
              ("25 个 Zernike 系数（单位 rad）。", dict(size=13.5))], space_after=6)
add_para(tf, [("• 物理本质：", dict(size=13.5, bold=True, color=INK)),
              ("用神经网络逼近 (E)NZ 反演算子，把昂贵的迭代优化替换为一次前向推理。", dict(size=13.5))], space_after=0)

# ----------------------------------------------------------------------------
# 19. Paper B - results
# ----------------------------------------------------------------------------
s = newslide("文献 B · 校正效果", "前沿文献 ②")
trip = [("rg1.png", "理想 (Ground truth)"),
        ("ra1.png", "含像差 (Aberrated)"),
        ("rp1.png", "校正后 (Corrected)")]
x0 = Inches(1.1); w = Inches(2.6); gap = Inches(0.5); top = Inches(1.7)
for i, (fn, lab) in enumerate(trip):
    lx = x0 + i * (w + gap)
    add_image_fit(s, f"{PB}/{fn}", lx, top, w, Inches(2.6))
    cb, cf = textbox(s, lx, top + Inches(2.6), w, Inches(0.35))
    add_para(cf, [(lab, dict(size=12, bold=True, color=INK))], first=True, align=PP_ALIGN.CENTER, space_after=0)
caption(s, Inches(0.7), Inches(4.75), Inches(12), "原文 Fig.4（PSF 数据集示例）：网络预测 Zernike 后重建/去除像差，PSF 基本恢复为艾里斑")
tb, tf = textbox(s, Inches(0.7), Inches(5.2), Inches(11.9), Inches(1.8))
add_para(tf, [("• 仿真 PSF：", dict(size=13.5, bold=True, color=ACC2)),
              ("3 张相位差异图像，离焦幅度=1 时 RMS≈0.10 rad；2 通道亦可，3 通道更稳。", dict(size=13.5))], first=True, space_after=6)
add_para(tf, [("• 扩展 2D 样品：", dict(size=13.5, bold=True, color=ACC2)),
              ("不再是点源，网络仍能直接预测 Zernike，RMS≈0.15 rad；可迭代 2~3 次进一步提升。", dict(size=13.5))], space_after=6)
add_para(tf, [("• 价值：", dict(size=14, bold=True, color=GREEN)),
              ("无需波前传感器、单步推理即可像差反演 → 离线“数字自适应光学”与超分辨成像的快速路径。",
               dict(size=13.5))], space_after=0)

# ----------------------------------------------------------------------------
# 20. Summary
# ----------------------------------------------------------------------------
s = newslide("总结与展望", "第三部分")
tb, tf = textbox(s, Inches(0.7), Inches(1.35), Inches(11.9), Inches(5.4))
add_para(tf, [("一条主线：", dict(size=17, bold=True, color=ACC)),
              ("波像差 W  →  光瞳相位 Φ  →  衍射积分  →  PSF / MTF  →  系统评价与反演。",
               dict(size=15))], first=True, space_after=12)
add_para(tf, [("基本原理小结：", dict(size=15, bold=True, color=ACC2))], space_after=4)
for t in [
    "用 Zernike 多项式正交展开像差；Nijboer–Zernike 给出逐项解析 PSF（贝塞尔函数）。",
    "Strehl 比 / Maréchal 判据 (S≥0.8 ⟺ σ_W≤λ/14) 给出“衍射极限”的定量标准。",
    "扩展 Nijboer–Zernike (ENZ) 把理论推广到大像差、离焦、复光瞳与高 NA 矢量衍射，并支持像差反演。",
]:
    add_para(tf, [("• " + t, dict(size=14))], space_after=5)
add_para(tf, [("前沿启示：", dict(size=15, bold=True, color=ACC)),], space_after=4)
add_para(tf, [("① 经典理论“活着”：", dict(size=14, bold=True, color=GREEN)),
              ("N–Z 模型被扩展用于空间引力波探测的远场波前误差与 TTL 噪声解析建模（文献 A）。",
               dict(size=14))], space_after=5)
add_para(tf, [("② 与 AI 融合：", dict(size=14, bold=True, color=GREEN)),
              ("深度网络学习 “PSF↔Zernike” 映射，把 (E)NZ 反问题变为一次前向推理，实现快速无传感器像差校正（文献 B）。",
               dict(size=14))], space_after=10)
rect(s, Inches(0.7), Inches(6.0), Inches(11.9), Inches(0.75), LIGHT)
tb2, tf2 = textbox(s, Inches(0.95), Inches(6.05), Inches(11.4), Inches(0.65), anchor=MSO_ANCHOR.MIDDLE)
add_para(tf2, [("一句话：", dict(size=14, bold=True, color=ACC)),
               ("像差的衍射理论既是百年经典，又在计算成像与人工智能时代焕发新生。", dict(size=14))],
         first=True, space_after=0)

# ----------------------------------------------------------------------------
# 21. References
# ----------------------------------------------------------------------------
s = newslide("参考文献 / References", "")
tb, tf = textbox(s, Inches(0.7), Inches(1.35), Inches(11.9), Inches(5.4))
refs = [
    "[1] M. Born & E. Wolf, Principles of Optics, Ch. 9 “The Diffraction Theory of Aberrations”, Cambridge Univ. Press.",
    "[2] B. R. A. Nijboer, “The diffraction theory of optical aberrations,” Physica 10 (1943) 679; 13 (1947) 605.",
    "[3] A. J. E. M. Janssen, “Extended Nijboer–Zernike approach to the computation of optical point-spread functions,” "
    "J. Opt. Soc. Am. A 19, 849 (2002); J. Braat et al., JOSA A 19, 858 (2002).",
    "[4] Y.-Z. Tao, R.-H. Gao, G. Xu, Y.-L. Wu, “Analytical Modeling of Far-Field Wavefront Error with Beam-Waist and "
    "Lateral-Shift Effects in Spaceborne Laser Interferometry,” arXiv:2604.26371 (2026). [CC BY 4.0]",
    "[5] Y. E. Kok, A. Bentley, A. Parkes, A. J. Wright, M. G. Somekh, M. Pound, “Direct Zernike Coefficient Prediction "
    "from Point Spread Functions and Extended Images using Deep Learning,” arXiv:2404.15231 (2024).",
]
for i, r in enumerate(refs):
    add_para(tf, [(r, dict(size=13.5))], first=(i == 0), space_after=10, line=1.1)
add_para(tf, [("图片说明：", dict(size=12, bold=True, color=GREY)),
              ("第一部分插图由本人用 Python (NumPy/SciPy/Matplotlib) 依据上述理论计算自绘；"
               "第二部分插图截取自对应文献 [4][5]（arXiv 公开预印本，仅用于课程学习，已标注来源）。",
               dict(size=12, color=GREY))], space_after=0, line=1.15)

# ----------------------------------------------------------------------------
out = f"{BASE}/像差的衍射理论_课程报告.pptx"
prs.save(out)
print("Saved:", out)
print("Slides:", len(prs.slides._sldIdLst))
